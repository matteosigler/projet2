"""Génère planning_arbre_vert.pptx : le planning de campagne « 25 ans de L'Arbre Vert »
reconstruit en formes et textes natifs, pour être importé et modifié dans Canva.

Les coordonnées sont exprimées en pixels de la maquette d'origine (1671 x 941).
"""
from datetime import date
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

HERE = Path(__file__).parent
OUT = HERE / "planning_arbre_vert.pptx"
LOGO = HERE / "logo_arbre_vert.png"

W_PX, H_PX = 1671, 941
K = 12192000 / W_PX  # EMU par pixel de maquette


def px(v):
    return Emu(int(round(v * K)))


def fpx(v):
    """Taille de police en pixels de maquette -> points."""
    return Pt(v * K / 12700)


# Palette
BG = "FCFAF6"
INK = "113B33"
INK_SOFT = "2A3A35"
GREY = "5F6B66"
DARK_BAR = "5B9360"
LIGHT_BAR = "AED0A7"
JOUR_J = "034034"
ACCENT = "238B5A"
HEADER_FILL = "EBF1E4"
PANEL_AMONT = "EAF1E2"
PANEL_TEMPS_FORT = "E1EEDA"
PANEL_APRES = "F2F3EC"
BAND_TEMPS_FORT = "EEF4E8"
GRID = "D6DBD2"
FOOTER_FILL = "EEF2E6"
LEGEND_FILL = "F3F4EC"

HEAD_FONT = "Montserrat"
BODY_FONT = "Lato"

# Échelle de temps : 1er août -> 31 décembre 2028
T0 = date(2028, 8, 1)
T_END = date(2029, 1, 1)
X0, X1 = 447, 1633
PPD = (X1 - X0) / (T_END - T0).days


def xd(d, frac=0.0):
    """Abscisse du début du jour d (+ fraction de jour)."""
    return X0 + ((d - T0).days + frac) * PPD


prs = Presentation()
prs.slide_width = px(W_PX)
prs.slide_height = px(H_PX)
slide = prs.slides.add_slide(prs.slide_layouts[6])


def rgb(hex_):
    return RGBColor.from_string(hex_)


def no_shadow(shape):
    sp_pr = shape._element.spPr
    if sp_pr.find(qn("a:effectLst")) is None:
        etree.SubElement(sp_pr, qn("a:effectLst"))
    style = shape._element.find(qn("p:style"))
    if style is not None:  # le style de thème ajoute une ombre portée
        shape._element.remove(style)


def box(x, y, w, h, fill, radius=None, name=None):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(kind, px(x), px(y), px(w), px(h))
    if radius:
        s.adjustments[0] = min(0.5, radius / min(w, h))
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(fill)
    s.line.fill.background()
    no_shadow(s)
    if name:
        s.name = name
    return s


def pill(x, y, w, h, fill, name=None):
    return box(x, y, w, h, fill, radius=h / 2, name=name)


def oval(cx, cy, r, fill, name=None):
    s = slide.shapes.add_shape(MSO_SHAPE.OVAL, px(cx - r), px(cy - r), px(2 * r), px(2 * r))
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(fill)
    s.line.fill.background()
    no_shadow(s)
    if name:
        s.name = name
    return s


def line(x1, y1, x2, y2, color, width, dash=None, name=None):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, px(x1), px(y1), px(x2), px(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = px(width)
    if dash:
        c.line.dash_style = dash
    no_shadow(c)
    if name:
        c.name = name
    return c


def text(x, y, w, h, runs, size, color=INK, bold=False, font=BODY_FONT,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE, spacing=None, name=None):
    """runs : str ou liste de (texte, {overrides}) pour du texte mixte."""
    tb = slide.shapes.add_textbox(px(x), px(y), px(w), px(h))
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    if isinstance(runs, str):
        runs = [(runs, {})]
    for content, over in runs:
        r = p.add_run()
        r.text = content
        f = r.font
        f.name = over.get("font", font)
        f.size = fpx(over.get("size", size))
        f.bold = over.get("bold", bold)
        f.color.rgb = rgb(over.get("color", color))
        sp = over.get("spacing", spacing)
        if sp is not None:
            r._r.get_or_add_rPr().set("spc", str(int(sp * 100)))
    if name:
        tb.name = name
    return tb


# ---------------------------------------------------------------- Fond
bg = slide.background.fill
bg.solid()
bg.fore_color.rgb = rgb(BG)

# ---------------------------------------------------------------- En-tête
text(48, 28, 1300, 74, "Planning de campagne · 25 ans de L'Arbre Vert",
     size=51, bold=True, font=HEAD_FONT, color=INK, spacing=-1.5, name="Titre")
text(50, 98, 1300, 38,
     "Annonce en amont  ·  temps fort le dimanche 24 septembre 2028 · prolongation jusqu'en décembre",
     size=26.5, color=GREY, name="Sous-titre")
slide.shapes.add_picture(str(LOGO), px(1530), px(19), px(91), px(104)).name = "Logo L'Arbre Vert"

# Légende
box(983, 142, 650, 56, LEGEND_FILL, radius=14, name="Légende fond")
pill(1009, 162, 45, 17, LIGHT_BAR, name="Légende préparation")
text(1066, 155, 200, 30, "Préparation / prolongation", size=17, color=INK_SOFT)
pill(1279, 162, 45, 17, DARK_BAR, name="Légende diffusion")
text(1336, 155, 170, 30, "Diffusion / temps fort", size=17, color=INK_SOFT)
line(1511, 152, 1511, 189, JOUR_J, 2, MSO_LINE_DASH_STYLE.DASH, name="Légende ligne Jour J")
oval(1539, 171, 9, JOUR_J, name="Légende point Jour J")
text(1563, 155, 60, 30, "Jour J", size=17, color=INK_SOFT)

# ---------------------------------------------------------------- Mois
months = [
    ("AOÛT 2028", date(2028, 8, 1), date(2028, 9, 1)),
    ("SEPTEMBRE 2028", date(2028, 9, 1), date(2028, 10, 1)),
    ("OCTOBRE 2028", date(2028, 10, 1), date(2028, 11, 1)),
    ("NOVEMBRE 2028", date(2028, 11, 1), date(2028, 12, 1)),
    ("DÉCEMBRE 2028", date(2028, 12, 1), T_END),
]
for label, start, end in months:
    x, xe = xd(start) + 1.5, xd(end) - 1.5
    box(x, 212, xe - x, 45, HEADER_FILL, radius=6, name=f"Mois {label}")
    text(x, 212, xe - x, 45, label, size=20, bold=True, font=HEAD_FONT,
         color=INK, align=PP_ALIGN.CENTER)

# ---------------------------------------------------------------- Panneaux de gauche
box(40, 243, 398, 322, PANEL_AMONT, radius=14, name="Panneau En amont")
box(40, 573, 398, 103, PANEL_TEMPS_FORT, radius=14, name="Panneau Temps fort")
box(447, 573, X1 - 447, 103, BAND_TEMPS_FORT, radius=10, name="Bande Temps fort")
box(40, 684, 398, 172, PANEL_APRES, radius=14, name="Panneau Après")

# Grille : débuts de mois + milieux de mois
grid_days = [date(2028, m, d) for m in (8, 9, 10, 11, 12) for d in (1, 16)] + [T_END]
for d in grid_days:
    x = xd(d)
    line(x, 264, x, 857, GRID, 1, MSO_LINE_DASH_STYLE.SQUARE_DOT, name="Grille")

# ---------------------------------------------------------------- En amont
text(70, 255, 360, 30, "EN AMONT", size=22, bold=True, font=HEAD_FONT, spacing=0.5)

TITLE_SIZE, SUB_SIZE, DATE_SIZE = 21, 17.5, 19
amont = [
    # (titre, sous-titre, centre y, début, fin, couleur)
    ("Posts clés réseaux sociaux", "stories · posts · carrousels · 5 clés / semaine", 315,
     date(2028, 8, 22), date(2028, 9, 23), DARK_BAR),
    ("Flyers dans les colis", None, 370, date(2028, 8, 22), date(2028, 9, 23), DARK_BAR),
    ("QR code sur les packagings", None, 424, date(2028, 8, 15), date(2028, 12, 31), LIGHT_BAR),
    ("Landing page", None, 476, date(2028, 8, 15), date(2028, 12, 31), LIGHT_BAR),
    ("Début de l'influence", None, 530, date(2028, 8, 22), date(2028, 9, 23), DARK_BAR),
]
MOIS = {8: "août", 9: "sept.", 10: "oct.", 11: "nov.", 12: "déc."}


def fr(d):
    return f"{d.day} {MOIS[d.month]}"


BAR_H = 32
for titre, sous_titre, cy, start, end, color in amont:
    if sous_titre:
        text(72, cy - 17, 360, 26, titre, size=TITLE_SIZE, bold=True, color=INK, name=titre)
        text(72, cy + 9, 360, 22, sous_titre, size=SUB_SIZE, color=GREY)
    else:
        text(72, cy - 14, 360, 28, titre, size=TITLE_SIZE, bold=True, color=INK, name=titre)
    xs, xe = xd(start), xd(end, 1)
    pill(xs, cy - BAR_H / 2, xe - xs, BAR_H, color, name=f"Barre {titre}")
    text(xs - 130, cy - 14, 121, 28, fr(start), size=DATE_SIZE, color=INK_SOFT, align=PP_ALIGN.RIGHT)
    if xe + 100 <= W_PX:
        text(xe + 14, cy - 14, 110, 28, fr(end), size=DATE_SIZE, color=INK_SOFT)
    else:  # barre jusqu'au bord : date de fin dans la barre
        text(xe - 130, cy - 14, 116, 28, fr(end), size=DATE_SIZE, color=INK_SOFT, align=PP_ALIGN.RIGHT)

# ---------------------------------------------------------------- Temps fort
text(72, 585, 360, 30, "SEPTEMBRE · TEMPS FORT", size=21.5, bold=True, font=HEAD_FONT, spacing=0.2)
text(72, 618, 360, 26, "Événement", size=21.5, bold=True, color="1F5C45")
text(72, 645, 360, 22, "Le Passage de la Clé", size=SUB_SIZE, color=GREY)

JJ = date(2028, 9, 24)
xj = xd(JJ, 0.5)

# ---------------------------------------------------------------- Après l'événement
text(70, 699, 360, 30, "APRÈS L'ÉVÉNEMENT", size=22, bold=True, font=HEAD_FONT, spacing=0.5)
apres = [
    # (titre, sous-titre, y du titre, centre y de la barre, début, fin)
    ("Influence continue + #PasseLaClé", "récupération des vidéos", 737, 742,
     date(2028, 9, 25), date(2028, 11, 19)),
    ("Création de nouvelles clés", "grâce au QR code et aux retours", 798, 807,
     date(2028, 9, 25), date(2028, 12, 17)),
]
for titre, sous_titre, ty, cy, start, end in apres:
    text(72, ty, 360, 26, titre, size=TITLE_SIZE, bold=True, color=INK, name=titre)
    text(72, ty + 26, 360, 22, sous_titre, size=SUB_SIZE, color=GREY)
    xs, xe = xd(start), xd(end, 1)
    pill(xs, cy - 17, xe - xs, 34, DARK_BAR, name=f"Barre {titre}")
    text(xs + 18, cy - 14, 110, 28, fr(start), size=DATE_SIZE, color="FFFFFF")
    text(xe + 10, cy - 14, 110, 28, fr(end), size=DATE_SIZE, color=INK_SOFT)

# ---------------------------------------------------------------- Ligne et point Jour J (au premier plan)
line(xj, 264, xj, 857, JOUR_J, 2, MSO_LINE_DASH_STYLE.DASH, name="Ligne Jour J")
oval(xj, 621, 15, JOUR_J, name="Point Jour J")
text(xj + 32, 601, 400, 28, "dimanche 24 septembre 2028", size=22, bold=True, color=INK)
text(xj + 32, 628, 200, 26, "Jour J", size=21.5, bold=True, color=ACCENT)

# ---------------------------------------------------------------- Pied de page
box(40, 870, 1593, 49, FOOTER_FILL, radius=14, name="Pied de page")
oval(91, 895, 24, "156449", name="Icône fond")
gear = slide.shapes.add_shape(MSO_SHAPE.GEAR_9, px(79), px(883), px(24), px(24))
gear.fill.background()
gear.line.color.rgb = rgb("FFFFFF")
gear.line.width = px(2)
no_shadow(gear)
gear.name = "Icône engrenage"
oval(91, 895, 4, "FFFFFF", name="Icône centre")
line(141, 883, 141, 908, INK, 2.5, name="Séparateur")
text(161, 878, 1450, 34, [
    ("Préparation interne · janv. → août : ", {"bold": True, "color": INK, "size": 20}),
    ("validation des 25 clés, contenus, visuels, organisation de l'événement.", {"color": INK_SOFT}),
], size=18.5, name="Préparation interne")

prs.save(OUT)
print(f"OK -> {OUT}")
